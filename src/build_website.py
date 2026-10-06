"""Build the complete website snapshot from a checked-out analytical repository."""
import argparse,csv,hashlib,json,shutil,sqlite3,subprocess,zipfile
from datetime import date,datetime
from pathlib import Path
import openpyxl
SITE=Path(__file__).resolve().parents[1]
PURPOSES={
 'How_to_read':'Workbook guidance and interpretation', 'Data_Dictionary':'Original field definitions and examples',
 'Data_Quality':'Quality summary for the unavailable full historical extract', 'KPI_Summary':'Original historical summaries and sample formulas',
 'Paid_Sample':'Selected 200-record historical high-fee sample · EUR', 'By_Position':'Original last-five-year position summary · EUR; a separate population',
 'By_Age':'Original age-band summary · EUR; a separate population', 'FTV_Worth':'Selected 150-player exploratory valuation sample · EUR',
 'Summer26_Overview':'Source coverage and summer methodology', 'Summer26_New':'Transfers newly added to the workbook · GBP quotes',
 'Summer26_Matches':'Summer source records already matched to the workbook', 'Summer26_Review':'Review/excluded records; excluded from summer dashboard',
 'Fee_Components':'Reported fee detail; never automatically added to quoted fees', 'Source_Checks':'Source verification notes and conflicting reports'}
NUMERIC={'quoted_fee_gbp','existing_excel_row','existing_fee_eur','transfer_fee','mv_at_transfer','fee_to_mv','age_at_transfer','age','tm_value','fair_value','tm_to_fair','minutes','goals','assists','amount','base_fee','addons_max','has_reported_fee','is_permanent_fee','editorial_worth_eur','effective_worth_eur','worth_rank'}
def scalar(v):return v.isoformat()[:10] if isinstance(v,(date,datetime)) else v
def csv_records(path):
 with path.open(encoding='utf8',newline='') as f:rows=list(csv.DictReader(f))
 for row in rows:
  for k,v in row.items():
   if v=='':row[k]=None
   elif k in NUMERIC:row[k]=float(v) if '.' in v else int(v)
 return rows
def pack(folder,output):
 with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
  for p in sorted(folder.rglob('*')):
   if p.is_file() and not any(x in {'.git','__pycache__','.pbi','downloads','.env'} for x in p.relative_to(folder).parts):z.write(p,folder.name+'/'+p.relative_to(folder).as_posix())
def main(source):
 dist=SITE/'website';downloads=dist/'downloads';downloads.mkdir(parents=True,exist_ok=True)
 workbook=source/'excel/Football_Transfer_Worth_Final.xlsx'
 values=openpyxl.load_workbook(workbook,data_only=True);formulas=openpyxl.load_workbook(workbook,data_only=False)
 sheets=[]
 for s in values:
  raw=[[scalar(c.value) for c in row] for row in s]
  structured=s.title not in {'How_to_read','Summer26_Overview','Position_Rubric'}
  header_row=5 if s.title in {'Transfer_Ratings','Rating_Calculator'} else 1
  if header_row==5:raw=raw[4:]
  headers=[str(v) if v is not None else openpyxl.utils.get_column_letter(i+1) for i,v in enumerate(raw[0])] if structured else [openpyxl.utils.get_column_letter(i+1) for i in range(s.max_column)]
  rows=raw[1:] if structured else raw
  if header_row==5:rows=[row for row in rows if any(v is not None for v in row)]
  fs=[{'cell':c.coordinate,'formula':c.value,'cached':scalar(s[c.coordinate].value)} for row in formulas[s.title] for c in row if c.data_type=='f']
  chart_specs=[]
  for chart in formulas[s.title]._charts:
   title=''
   if chart.title and chart.title.tx and chart.title.tx.rich:
    title=''.join(r.t for p in chart.title.tx.rich.p for r in p.r)
   chart_specs.append({'type':type(chart).__name__,'title':title,'series':[{'values':getattr(getattr(x.val,'numRef',None),'f',None)} for x in chart.series]})
  sheets.append({'name':s.title,'purpose':PURPOSES.get(s.title,'Ratings and position-scoring framework'),'headerRow':header_row,'structured':structured,'columns':headers,'rows':rows,'sheetRowCount':s.max_row,'sheetColumnCount':s.max_column,'formulas':fs,'charts':chart_specs,'tables':[{'name':t.name,'range':t.ref} for t in formulas[s.title].tables.values()]})
  with (downloads/f'{s.title}.csv').open('w',newline='',encoding='utf-8-sig') as f:
   writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)
 datasets={p.stem:csv_records(p) for p in (source/'data/processed').glob('*.csv')}
 summer=datasets['summer_transfers'];ids=[r['record_id'] for r in summer]
 assert len(summer)==1493 and len(ids)==len(set(ids))
 assert sum(r['quoted_fee_gbp'] is not None and r['quoted_fee_gbp']>0 for r in summer)==499
 mapping={'summer_new':'Summer26_New','summer_matches':'Summer26_Matches','summer_review':'Summer26_Review','paid_sample':'Paid_Sample','valuation_sample':'FTV_Worth','fee_components':'Fee_Components','source_checks':'Source_Checks'}
 for dataset,sheet in mapping.items():
  obj=next(x for x in sheets if x['name']==sheet)
  expected=[dict(zip(obj['columns'],row)) for row in obj['rows'] if any(v is not None for v in row)]
  assert len(expected)==len(datasets[dataset]),dataset
  for a,b in zip(expected,datasets[dataset]):
   assert all(a[k]==b.get(k) for k in a),dataset
 with sqlite3.connect(source/'data/transfer_market.sqlite') as con:
  total=con.execute('SELECT SUM(quoted_fee_gbp) FROM summer_transfers WHERE is_permanent_fee=1').fetchone()[0]
  assert abs(total-sum(r['quoted_fee_gbp'] or 0 for r in summer if r['is_permanent_fee']))<.01
  con.row_factory=sqlite3.Row
  results={p.stem:[dict(r) for r in con.execute(p.read_text(encoding='utf8'))] for p in sorted((source/'sql/analysis').glob('*.sql'))}
 commit=subprocess.check_output(['git','-C',str(source),'rev-parse','HEAD'],text=True).strip()
 sources={'commit':commit,'workbookHash':hashlib.sha256(workbook.read_bytes()).hexdigest(),'snapshotDate':'2026-10-03','repository':'https://github.com/CopiousHighlights/football-transfer-market-analysis'}
 report=json.loads((source/'analysis/quality_report.json').read_text(encoding='utf8'))
 queries={p.name:p.read_text(encoding='utf8') for p in sorted((source/'sql/analysis').glob('*.sql'))}
 pbi={'measures':(source/'power-bi/measures.dax').read_text(encoding='utf8'),'visuals':[json.loads(p.read_text(encoding='utf8')) for p in (source/'power-bi').rglob('visual.json')],'status':'Schema checked; Desktop visual rendering, filters and DAX execution remain unverified.'}
 out={'sources':sources,'sheets':sheets,'datasets':datasets,'sqlResults':results,'queries':queries,'quality':report,'powerBI':pbi,'pythonExcerpt':(source/'src/pipeline.py').read_text(encoding='utf8')[:7000],'methodology':(source/'docs/METHODOLOGY.md').read_text(encoding='utf8')}
 (dist/'data.json').write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')),encoding='utf8')
 shutil.copy2(workbook,downloads/workbook.name);shutil.copy2(source/'data/transfer_market.sqlite',downloads/'transfer_market.sqlite')
 pack(source/'power-bi',downloads/'PowerBI_Complete_Project.zip');pack(source,downloads/'Football_Transfer_Complete_Project.zip')
 with zipfile.ZipFile(downloads/'Processed_Datasets.zip','w',zipfile.ZIP_DEFLATED) as z:
  for p in (source/'data/processed').glob('*.csv'):z.write(p,p.name)
 with zipfile.ZipFile(downloads/'SQL_Queries_and_Results.zip','w',zipfile.ZIP_DEFLATED) as z:
  for folder in ['sql','analysis']:
   for p in (source/folder).rglob('*'):
    if p.is_file():z.write(p,p.relative_to(source).as_posix())
 shutil.copy2(source/'docs/METHODOLOGY.md',downloads/'METHODOLOGY.md')
 shutil.copy2(source/'data/processed/ratings.json',dist/'ratings.json')
 shutil.copy2(source/'data/processed/transfer_ratings.csv',downloads/'CDM_Transfer_Ratings.csv')
 shutil.copy2(source/'docs/RATING_METHODOLOGY.md',downloads/'CDM_Rating_Methodology.md')
 audit={'commit':commit,'sheetCounts':{s['name']:len(s['rows']) for s in sheets},'summerRecords':len(summer),'positiveFeeRecords':499,'reportedPermanentFeesGBP':total,'workbookCSVReconciliation':'passed','SQLFeeReconciliation':'passed','powerBIVisualDefinitions':len(pbi['visuals'])}
 (SITE/'data-audit.json').write_text(json.dumps(audit,indent=2),encoding='utf8')
 print(json.dumps(audit,indent=2))
if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,default=SITE);main(parser.parse_args().source.resolve())
