import json,re,unittest
from pathlib import Path
from zipfile import ZipFile
import openpyxl
R=Path(__file__).resolve().parents[1]
class RatingTests(unittest.TestCase):
    def test_transfer_identity_and_missing_fees(self):
        rows=json.loads((R/'data/processed/ratings.json').read_text(encoding='utf-8'))['rows']
        self.assertEqual(len(rows),len({r['rating_id'] for r in rows}))
        self.assertEqual(len(rows),711)
        strikers=[r for r in rows if r['dataset']=='historical' and r['role'].startswith('Striker')]
        self.assertEqual(len(strikers),55)
        agreed={'Alexander Isak':7.1,'Hugo Ekitiké':8.8,'Erling Haaland':9.8,'Harry Kane':9.5,'Dominic Solanke':6.7}
        for player,score in agreed.items():
            references=[r for r in strikers if r['player_name']==player and r['rating_status']=='Your rating']
            self.assertEqual(len(references),1)
            self.assertEqual(references[0]['rating'],score)
        pending=[r for r in rows if r['rating_status']=='Evidence needed']
        self.assertEqual(len(pending),504)
        self.assertTrue(all(r['rating'] is None for r in pending))
        wingers=[r for r in rows if r['dataset']=='historical' and r['role']=='Winger']
        self.assertEqual(len(wingers),45)
        self.assertEqual(sum(r['rating'] is not None for r in wingers),40)
        references={('Raphinha','Barcelona'):9.7,('Michael Olise','Bayern Munich'):9.8,('Luis Díaz','Liverpool'):9.4,('Luis Díaz','Bayern Munich'):8.8,('Jérémy Doku','Man City'):9.1,('Ousmane Dembélé','PSG'):9.3,('Jadon Sancho','Man Utd'):2.4,('Khvicha Kvaratskhelia','PSG'):8.6,('Antony','Man Utd'):2.2}
        for (player,buyer),score in references.items():
            matching=[r for r in wingers if r['player_name']==player and r['to_club']==buyer]
            self.assertEqual(len(matching),1)
            self.assertEqual(matching[0]['rating'],score)
        for row in rows:
            if row['rating'] is not None:self.assertTrue(.1<=row['rating']<=10)
            if row['rating_status']=='Fee needed':self.assertIsNone(row['rating'])
        rodri=[r for r in rows if r['dataset']=='historical' and r['player_name']=='Rodri']
        self.assertEqual({r['rating'] for r in rodri},{8.8,9.8})
    def test_original_worksheets_and_native_parts_preserved(self):
        changed={'xl/styles.xml','xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','docProps/app.xml','xl/worksheets/sheet8.xml'}
        with ZipFile(R/'excel/original/Football_Transfer_Worth_Final_Original.xlsx') as old,ZipFile(R/'excel/Football_Transfer_Worth_Final.xlsx') as current:
            for name in old.namelist():
                if name not in changed:self.assertEqual(old.read(name),current.read(name),name)
    def test_editorial_worth_preserves_model_and_ranks_yamal_first(self):
        w=openpyxl.load_workbook(R/'excel/Football_Transfer_Worth_Final.xlsx',data_only=True)
        old=openpyxl.load_workbook(R/'excel/original/Football_Transfer_Worth_Final_Original.xlsx',data_only=True)
        current=w['FTV_Worth'];source=old['FTV_Worth']
        expected={'Lamine Yamal':220000000,'Erling Haaland':195000000,'Antony':95000000}
        for row in range(1,152):
            self.assertEqual([current.cell(row,c).value for c in range(1,13)],[source.cell(row,c).value for c in range(1,13)])
            player=current.cell(row,1).value
            if player in expected:
                self.assertEqual(current.cell(row,13).value,expected[player])
                self.assertEqual(current.cell(row,14).value,expected[player])
                self.assertEqual(current.cell(row,15).value,'Owner editorial estimate')
                if player=='Lamine Yamal':self.assertEqual(current.cell(row,16).value,1)
    def test_position_weights_and_formula_caches(self):
        w=openpyxl.load_workbook(R/'excel/Football_Transfer_Worth_Final.xlsx',data_only=True)
        for values in w['Position_Rubric'].iter_rows(min_row=6,max_row=14,values_only=True):self.assertAlmostEqual(sum(values[2:7]),1)
        self.assertIsNone(w['Rating_Calculator']['J6'].value)
        self.assertEqual(w['Rating_Calculator']['L6'].value,'Needs component scores')
        self.assertEqual(w['Transfer_Ratings']['K9'].value,9.7)
    def test_excel_compatibility_namespace_declarations(self):
        with ZipFile(R/'excel/Football_Transfer_Worth_Final.xlsx') as z:
            for part in ['xl/workbook.xml','xl/styles.xml']:
                text=z.read(part).decode('utf-8')
                prefixes=set(re.findall(r'xmlns:([\w]+)=',text))
                for tokens in re.findall(r'(?:Ignorable|Requires)="([^"]+)"',text):
                    self.assertTrue(set(tokens.split()).issubset(prefixes),part)
if __name__=='__main__':unittest.main()
